"""
Generate a PowerPoint presentation (PPTX) from the project outline.

Requirements:
  - Python 3.8+
  - python-pptx (install with: pip install python-pptx)

Run:
  python tools/generate_presentation.py

This script will create: BillingInvoicingPlatform_Presentation.pptx
in the repository root.
"""
from pptx import Presentation
from pptx.util import Inches, Pt

slides = [
    {
        "title": "Billing & Invoicing Platform (Enterprise)",
        "bullets": [
            "Backend API for invoice lifecycle, payments, reports and email delivery",
            ".NET 8, C# 12, EF Core, ASP.NET Core Identity, JWT, Hangfire, QuestPDF, MailKit",
            "AutoMapper, FluentValidation, Swagger, Rate Limiting"
        ],
        "notes": "Introduce the repository: a backend API implementing core billing features — customers, invoices, payments, reporting, email/PDF generation and background jobs. Emphasize API-first approach and production-grade libraries.",
        "diagram": "Stack: API -> SQL Server, Identity, Hangfire, SMTP, PDF service"
    },
    {
        "title": "Business Perspective",
        "bullets": [
            "What: backend system to create, send and manage invoices and payments",
            "Goal: automate invoice lifecycle, notifications (email + PDF), and reporting",
            "Target users: SMEs, accountants, internal finance teams, freelancers",
            "Value: workflow enforcement, automation, centralized reporting"
        ],
        "notes": "Explain what the application is and the main goal. Emphasize reduced manual work and value for SMEs and finance teams. Mention that repository is API-only (no UI included).",
        "diagram": "User personas -> API"
    },
    {
        "title": "Problem Statement",
        "bullets": [
            "Manual invoicing leads to missed payments and inconsistent tracking",
            "Need for role-based secure API for billing operations",
            "Need automated delivery (PDF + email) and overdue detection"
        ],
        "notes": "Map each problem to implemented features: JWT + Identity for security, Hangfire for background tasks, QuestPDF + MailKit for PDF and email.",
        "diagram": "Problems -> Solution components"
    },
    {
        "title": "Project Overview",
        "bullets": [
            "Customer CRUD with pagination and validation",
            "Invoice creation, update, send, cancel, delete with business rules",
            "Record/delete payments with transactional integrity",
            "Reports: outstanding receivables and revenue summaries",
            "Email delivery with PDF attachment via background jobs"
        ],
        "notes": "Mention controllers: Customers, Invoices, Payments, Reports, Accounts for auth. Describe capabilities in one minute overview.",
        "diagram": "Controllers -> Services -> Repositories"
    },
    {
        "title": "Architecture Overview",
        "bullets": [
            "API layer: ASP.NET Core controllers (REST)",
            "Application layer: Services containing business rules",
            "Infrastructure: Repositories (EF Core) and DbContext",
            "Cross-cutting: AutoMapper, FluentValidation, Exception middleware",
            "Background: Hangfire jobs for email and overdue processing"
        ],
        "notes": "Explain responsibilities of each layer and the separation of concerns. Note repository projections and service-level rules.",
        "diagram": "Layered architecture diagram (API -> Services -> Repos -> DB)"
    },
    {
        "title": "Why this Architecture",
        "bullets": [
            "Separation of concerns for maintainability and testability",
            "Repositories enable optimized SQL projections",
            "Services centralize domain/business rules",
            "Hangfire provides reliable background processing with dashboard"
        ],
        "notes": "Explain pragmatic Clean-like architecture and why it helps in this project. Mention code evidence (projections in InvoiceRepository, services for rules).",
        "diagram": "Rationale arrows linking layers to benefits"
    },
    {
        "title": "Technology Stack",
        "bullets": [
            ".NET 8, C# 12",
            "ASP.NET Core Web API, EF Core (SQL Server)",
            "ASP.NET Core Identity + JWT, Hangfire (SQL storage)",
            "QuestPDF (PDF), MailKit (SMTP), AutoMapper, FluentValidation, Swagger"
        ],
        "notes": "Point to Program.cs for DI registrations and configuration. Highlight use of these libraries in code.",
        "diagram": "Stack list"
    },
    {
        "title": "Authentication & Security",
        "bullets": [
            "JWT flow via AuthService (Register/Login produce token)",
            "Token validation configured in JwtBearer (issuer, audience, signing key)",
            "Role-based authorization attributes on controllers/actions",
            "Identity options: password complexity, lockout, unique email",
            "Rate limiting and HTTPS enforced, custom exception middleware"
        ],
        "notes": "Explain CreateJwtToken in AuthService, role claims, ClockSkew set to zero, and controller-level Authorize usage.",
        "diagram": "Auth flow: Login -> JWT -> Access endpoints"
    },
    {
        "title": "Customers Module",
        "bullets": [
            "Endpoints secured and rate-limited (CustomersController)",
            "CustomerService handles CRUD, validation and pagination",
            "CustomerRepository uses EF Core; AutoMapper profile maps Address value object",
            "Pagination metadata added to response headers (X-Pagination)"
        ],
        "notes": "Describe endpoint roles: Admin/Accountant for create/update, Admin for delete. Mention rate limiting. Refer to CustomerProfile mapping.",
        "diagram": "Customers flow: Controller -> Service -> Repository"
    },
    {
        "title": "Invoices Lifecycle",
        "bullets": [
            "Create invoice (Draft) with items; totals computed in service",
            "Update only allowed for Draft invoices",
            "Send changes status to Sent and enqueues email job",
            "Cancel rules: cannot cancel Paid/PartiallyPaid or invoices with payments",
            "Invoice numbering: INV-{Year}-{sequence:D5}"
        ],
        "notes": "Walk through InvoiceService methods: CreateAsync, UpdateAsync, SendInvoiceAsync, CancelInvoiceAsync and business validations. Highlight background enqueue for emails.",
        "diagram": "Draft -> Sent -> PartiallyPaid/Paid/Overdue transitions"
    },
    {
        "title": "Payments Handling",
        "bullets": [
            "Record payment with transactional UnitOfWork",
            "Validations: cannot pay Draft/Cancelled/Paid invoices",
            "Payment amount checked against remaining balance",
            "Delete payment within transaction and recalculate invoice status",
            "Invoice status determined by remaining balance (Paid or PartiallyPaid)"
        ],
        "notes": "Point to PaymentService: RecordPayment and DeletePaymentAsync with BeginTransaction/Commit/Rollback flows. Describe DetermineInvoiceStatus behaviour.",
        "diagram": "Payment -> Invoice status recalculation"
    },
    {
        "title": "Reports Module",
        "bullets": [
            "Outstanding receivables report with filters and summary statistics",
            "Revenue summary report with customer breakdown and optional comparison",
            "ReportService validates query and computes summaries",
            "InvoiceRepository provides optimized SQL projections for reports"
        ],
        "notes": "Explain ReportService and how repository methods return DTOs with SQL-side calculations (RemainingBalance, DaysOverdue).",
        "diagram": "Reports query -> repository projections -> summary"
    },
    {
        "title": "Email, PDF & Background Jobs",
        "bullets": [
            "PDF generation with QuestPDF (InvoicePdfService)",
            "HTML email and SMTP sending with MailKit (EmailService)",
            "InvoiceEmailJob orchestrates PDF generation and sending",
            "InvoiceOverdueJob runs daily to update invoice statuses (Hangfire recurring job)",
            "Hangfire dashboard enabled at /hangfireDashboard"
        ],
        "notes": "Describe email HTML body and attachment flow in EmailService and InvoiceEmailJob sequence. Note logging and retry via Hangfire.",
        "diagram": "Hangfire -> InvoiceEmailJob -> EmailService/InvoicePdfService -> SMTP"
    },
    {
        "title": "API Flow Scenario (Demo Script)",
        "bullets": [
            "Auth: POST /api/accounts/login -> receive JWT",
            "Customer: POST /api/customers (Admin/Accountant)",
            "Invoice: POST /api/invoices -> Draft, then PATCH /status -> Sent (enqueues email)",
            "Payment: POST /api/payments -> record payment and update invoice status",
            "Reporting: GET /api/reports/outstanding-receivables and /revenue-summary"
        ],
        "notes": "Step through endpoints and which services/controllers are involved at each step. Mention background job timing and Hangfire dashboard for monitoring.",
        "diagram": "Sequence diagram: Auth -> Customer -> Invoice -> Payment -> Report"
    },
    {
        "title": "Technical Challenges & Solutions",
        "bullets": [
            "Performance: repository projections and SQL calculations to reduce memory usage",
            "Validation: FluentValidation and service-level business checks",
            "Error handling: centralized ExceptionMiddleWare",
            "Consistency: UnitOfWork usage for payment operations",
            "Enum-to-string search limitation noted in repository code comments"
        ],
        "notes": "Explain each challenge and cite code locations: InvoiceRepository projections, PaymentService transactions, ExceptionMiddleWare registration in Program.cs. Mention the TODO regarding enum string search.",
        "diagram": "Problem -> Implemented solution mapping"
    },
    {
        "title": "Key Strengths",
        "bullets": [
            "Clear separation of concerns and service-oriented domain logic",
            "Background processing and durable jobs via Hangfire",
            "Secure Identity + JWT with role-based access",
            "Optimized data access patterns and modular external services"
        ],
        "notes": "Summarize why this codebase is close to production-ready and highlight practical strengths such as logging, Swagger and DI configuration.",
        "diagram": "Checklist-style strengths"
    },
    {
        "title": "Known Gaps & Unclear Areas",
        "bullets": [
            "No payment gateway integration (only internal recording)",
            "Seeding code commented out in Program.cs (roles/users not automatically created)",
            "No visible unit/integration tests in repository",
            "No front-end/UI bundled with this repository"
        ],
        "notes": "Be explicit these are facts derived from the repository contents. Encourage adding tests and gateway integration when needed.",
        "diagram": "Gaps -> Next steps"
    },
    {
        "title": "Future Improvements",
        "bullets": [
            "Add unit and integration tests for services and repositories",
            "Implement status-text search or map enum to DB string for search",
            "Add optimistic concurrency (rowversion) and more robust seeding",
            "Integrate payment gateways (Stripe/PayPal) and monitoring/metrics"
        ],
        "notes": "Link each improvement to where it integrates in current codebase (tests around InvoiceService/PaymentService, monitoring in Program.cs).",
        "diagram": "Roadmap timeline"
    },
    {
        "title": "Deployment & Operational Notes",
        "bullets": [
            "SQL Server is used for both application data and Hangfire storage",
            "SMTP credentials required in MailSettings for email delivery",
            "Recurring job configured in Program.cs (daily overdue run)",
            "Swagger enabled for development testing"
        ],
        "notes": "Mention the Hangfire dashboard protection and that appsettings must include JWT and mail settings for production.",
        "diagram": "Ops checklist"
    },
    {
        "title": "Conclusion",
        "bullets": [
            "Mature backend for billing operations with secure access and background processing",
            "Service and repository separation aids maintainability and extensibility",
            "Recommended next steps: tests, payment integrations, monitoring and front-end"
        ],
        "notes": "Provide a concise closing summary and invite further questions or follow-up tasks.",
        "diagram": "Summary box"
    },
    {
        "title": "Appendix: Key Code References",
        "bullets": [
            "Auth: Infrastructure/Identity/Service/AuthService.cs",
            "Invoices: Application/Service/InvoiceService.cs",
            "Payments: Application/Service/PaymentService.cs",
            "Repositories: Infrastructure/Repositories/InvoiceRepository.cs",
            "Background jobs: Infrastructure/BackgroundJobs/*.cs",
            "Startup: API/Program.cs"
        ],
        "notes": "Point reviewers to these files for deeper reading of business rules and integrations.",
        "diagram": "File tree excerpt"
    }
]

prs = Presentation()
blank_slide_layout = prs.slide_layouts[6]
# Title slide layout: try layout 0 if available
try:
    title_layout = prs.slide_layouts[0]
except Exception:
    title_layout = blank_slide_layout

for idx, s in enumerate(slides):
    # Use title layout for first slide
    if idx == 0:
        slide = prs.slides.add_slide(title_layout)
        if slide.shapes.title:
            slide.shapes.title.text = s['title']
        # Add subtitle as small textbox
        left = Inches(1)
        top = Inches(1.9)
        width = Inches(8)
        height = Inches(1.5)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        for i, b in enumerate(s['bullets']):
            p = tf.add_paragraph() if i>0 or tf.text else tf.paragraphs[0]
            p.text = b
            p.font.size = Pt(18)
    else:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title:
            slide.shapes.title.text = s['title']
        # content placeholder is typically at index 1
        try:
            body = slide.shapes.placeholders[1].text_frame
            body.clear()
        except Exception:
            # fallback textbox
            left = Inches(0.5)
            top = Inches(1.6)
            width = Inches(9)
            height = Inches(4.5)
            body_shape = slide.shapes.add_textbox(left, top, width, height)
            body = body_shape.text_frame
        # Add bullets
        for i, b in enumerate(s['bullets']):
            if i == 0:
                p = body.paragraphs[0]
                p.text = b
            else:
                p = body.add_paragraph()
                p.text = b
            p.level = 0
            p.font.size = Pt(14)

    # Add speaker notes
    notes_slide = slide.notes_slide
    notes_tf = notes_slide.notes_text_frame
    notes_text = s.get('notes','')
    diagram = s.get('diagram','')
    full_notes = notes_text
    if diagram:
        full_notes += "\n\nDiagram suggestion: " + diagram
    notes_tf.text = full_notes

output_path = "BillingInvoicingPlatform_Presentation.pptx"
prs.save(output_path)
print(f"Presentation generated: {output_path}")
